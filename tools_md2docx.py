# -*- coding: utf-8 -*-
"""Markdown -> Word (.docx) cho tai lieu du an.
   Giu: heading, BANG THAT (khong phai text), khoi code canh cot, **dam**, `code`, trich dan, list.
   Dung: python md2docx.py <input.md> <output.docx> [tieu de]"""
import io, os, re, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

MONO, BODY = 'Consolas', 'Calibri'
CODE_BG, CODE_PT, BODY_PT = 'F2F2F2', 8, 10.5

def _shade(el, hexcolor):
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), hexcolor)
    el.append(sh)

def _set_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    # font chu Dong A phai set rieng, khong thi Word doi font cho tieng Viet
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if size  is not None: run.font.size = Pt(size)
    if bold  is not None: run.font.bold = bold
    if color is not None: run.font.color.rgb = color

INLINE = re.compile(r'(\*\*.+?\*\*|`[^`]+`)')

def add_inline(par, text, base=BODY, size=BODY_PT, bold=False):
    """Tach **dam** va `code` thanh cac run rieng."""
    for tok in INLINE.split(text):
        if not tok: continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            _set_font(par.add_run(tok[2:-2]), base, size, True)
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 2:
            r = par.add_run(tok[1:-1])
            _set_font(r, MONO, size - 1.5, bold, RGBColor(0xB0, 0x30, 0x30))
        else:
            _set_font(par.add_run(tok), base, size, bold)

def add_code(doc, lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.18); pf.space_before = Pt(4); pf.space_after = Pt(8)
    pf.line_spacing = 1.0
    _shade(p._p.get_or_add_pPr(), CODE_BG)
    _set_font(p.add_run('\n'.join(lines)), MONO, CODE_PT)

def split_row(line):
    s = line.strip()
    if s.startswith('|'): s = s[1:]
    if s.endswith('|'):   s = s[:-1]
    return [c.strip() for c in s.split('|')]

def add_table(doc, rows):
    head, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(head))
    t.style = 'Table Grid'; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, c in enumerate(head):
        cell = t.rows[0].cells[i]
        cell.paragraphs[0].text = ''
        add_inline(cell.paragraphs[0], c, size=BODY_PT - 1, bold=True)
        _shade(cell._tc.get_or_add_tcPr(), 'E8EEF7')
    for r in body:
        cells = t.add_row().cells
        for i in range(len(head)):
            cells[i].paragraphs[0].text = ''
            add_inline(cells[i].paragraphs[0], r[i] if i < len(r) else '', size=BODY_PT - 1)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def convert(src, dst, title=None):
    md = io.open(src, encoding='utf-8').read().splitlines()
    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)          # A4
    sec.left_margin = sec.right_margin = Inches(0.7)
    sec.top_margin  = sec.bottom_margin = Inches(0.8)
    st = doc.styles['Normal']
    st.font.name = BODY; st.font.size = Pt(BODY_PT)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), BODY)

    if title:
        h = doc.add_heading(title, 0)
        for r in h.runs: _set_font(r, BODY, 22, True)

    i, n = 0, len(md)
    while i < n:
        line = md[i]

        if line.startswith('```'):                                  # khoi code
            j = i + 1
            while j < n and not md[j].startswith('```'): j += 1
            add_code(doc, md[i+1:j]); i = j + 1; continue

        if line.strip().startswith('|') and i + 1 < n and re.match(r'^\s*\|[\s:|-]+\|\s*$', md[i+1]):
            rows, j = [split_row(line)], i + 2       # bang: bo qua dong ngan cach
            while j < n and md[j].strip().startswith('|'):
                rows.append(split_row(md[j])); j += 1
            add_table(doc, rows); i = j; continue

        if re.match(r'^\s*(---|===)\s*$', line):                    # duong ke ngang
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
            pb = OxmlElement('w:pBdr'); bt = OxmlElement('w:bottom')
            bt.set(qn('w:val'), 'single'); bt.set(qn('w:sz'), '6'); bt.set(qn('w:color'), 'BBBBBB')
            pb.append(bt); p._p.get_or_add_pPr().append(pb); i += 1; continue

        m = re.match(r'^(#{1,4})\s+(.*)$', line)                    # heading
        if m:
            lvl, txt = len(m.group(1)), m.group(2)
            h = doc.add_heading('', min(lvl, 4))
            add_inline(h, txt, size={1: 17, 2: 14, 3: 12, 4: 11}[min(lvl, 4)], bold=True)
            for r in h.runs: r.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
            i += 1; continue

        if line.startswith('>'):                                    # trich dan
            buf, j = [], i
            while j < n and md[j].startswith('>'):
                buf.append(md[j].lstrip('>').strip()); j += 1
            p = doc.add_paragraph(); pf = p.paragraph_format
            pf.left_indent = Inches(0.3); pf.space_before = Pt(4); pf.space_after = Pt(6)
            _shade(p._p.get_or_add_pPr(), 'FFF8E1')
            add_inline(p, ' '.join(x for x in buf if x))
            i = j; continue

        m = re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)$', line)           # list
        if m:
            deep = len(m.group(1)) >= 2
            style = 'List Bullet' if m.group(2) in '-*' else 'List Number'
            p = doc.add_paragraph(style=style + (' 2' if deep else ''))
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, m.group(3)); i += 1; continue

        if not line.strip(): i += 1; continue

        buf, j = [], i                                              # doan van
        while j < n and md[j].strip() and not re.match(
                r'^(#{1,4}\s|```|\s*\||>|\s*([-*]|\d+\.)\s|---\s*$)', md[j]):
            buf.append(md[j].strip()); j += 1
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(6)
        add_inline(p, ' '.join(buf)); i = j

    doc.save(dst)
    return dst

if __name__ == '__main__':
    src, dst = sys.argv[1], sys.argv[2]
    title = sys.argv[3] if len(sys.argv) > 3 else None
    convert(src, dst, title)
    print(f'{dst}  ({os.path.getsize(dst)/1024:.0f} KB)')
